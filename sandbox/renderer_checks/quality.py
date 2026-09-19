"""Preserved renderer regression checks, independent of the retired Broker."""

import base64
import json
import os
import sys
from io import BytesIO
from zipfile import ZipFile

try:
    from ..job import execute as execute_request
except ImportError:
    from job import execute as execute_request


def check_builtin_skill_quality_guards() -> None:
    long_text = "这是一段用于验证长文本自动换行和一次生成成功的中文内容" * 5
    document = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "content": (
                        "# 中文报告\n\n这是中文内容。\n\n"
                        "| 项目 | 说明 | 负责人 | 状态 |\n"
                        "| --- | --- | --- | --- |\n"
                        f"| 交付 | {long_text} | 团队 | 完成 |"
                    )
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "quality.docx"},
        }
    )
    assert document["ok"] is True, document
    with ZipFile(
        BytesIO(base64.b64decode(document["artifact"]["content_base64"]))
    ) as archive:
        document_xml = archive.read("word/document.xml").decode()
        report_styles_xml = archive.read("word/styles.xml").decode()
        report_footer_xml = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/footer")
        )
    if not os.environ.get("NEXAFLOW_CJK_FONT"):
        report_cjk_font = (
            "Noto Serif CJK SC" if sys.platform == "linux" else "Songti SC"
        )
        assert report_cjk_font in report_styles_xml
        assert "Microsoft YaHei" not in document_xml
    assert '<w:tblW w:type="dxa" w:w="9071"' in document_xml
    assert '<w:tblLayout w:type="fixed"' in document_xml
    assert "1F4E79" in report_styles_xml
    assert b" PAGE " in report_footer_xml

    formal = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "style": "formal_legal",
                    "content": (
                        "# 劳动仲裁申请书\n\n"
                        "## 一、仲裁请求\n\n"
                        "1. 请求被申请人支付工资。\n\n"
                        "- 证据清单保持字面项目符号。\n\n"
                        "这是用于验证正式文书正文格式的段落。\n\n"
                        "---\n\n"
                        "## 附件说明\n\n"
                        "> 申请人：张三\n"
                        "> 2026年9月4日"
                    ),
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "formal-legal.docx"},
        }
    )
    assert formal["ok"] is True, formal
    formal_bytes = base64.b64decode(formal["artifact"]["content_base64"])
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    formal_document = Document(BytesIO(formal_bytes))
    paragraphs = {paragraph.text: paragraph for paragraph in formal_document.paragraphs}
    title = paragraphs["劳动仲裁申请书"]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.style.name == "Normal"
    assert all(str(run.font.color.rgb) == "000000" for run in title.runs)
    assert all(run.font.size.pt == 18 for run in title.runs)
    section_heading = paragraphs["一、仲裁请求"]
    assert section_heading.style.name == "Normal"
    assert all(str(run.font.color.rgb) == "000000" for run in section_heading.runs)
    assert all(run.font.size.pt == 14 for run in section_heading.runs)
    normal = formal_document.styles["Normal"]
    assert normal.font.size.pt == 12
    assert abs(normal.paragraph_format.line_spacing - 1.4) < 0.01
    assert abs(normal.paragraph_format.first_line_indent.pt - 24) < 0.1
    numbered = paragraphs["1. 请求被申请人支付工资。"]
    assert numbered.style.name == "Normal"
    bullet = paragraphs["- 证据清单保持字面项目符号。"]
    assert bullet.style.name == "Normal"
    assert all(
        paragraph.style.name == "Normal"
        for paragraph in formal_document.paragraphs
    )
    assert paragraphs["申请人：张三"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraphs["2026年9月4日"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    with ZipFile(BytesIO(formal_bytes)) as archive:
        formal_document_xml = archive.read("word/document.xml").decode()
        formal_styles_xml = archive.read("word/styles.xml").decode()
        footer_xml = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/footer")
        )
    formal_cjk_font = os.environ.get("NEXAFLOW_CJK_FONT") or (
        "Noto Serif CJK SC" if sys.platform == "linux" else "Songti SC"
    )
    assert formal_cjk_font in formal_styles_xml
    assert "---" not in formal_document_xml
    assert "&gt;" not in formal_document_xml
    assert '<w:br w:type="page"' in formal_document_xml
    assert b" PAGE " not in footer_xml

    conflicting = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "style": "formal_legal",
                    "content": (
                        "# 劳动仲裁申请书\n\n"
                        "仲裁费用及其他费用由被申请人承担。\n\n"
                        "劳动仲裁不收费。"
                    ),
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "conflict.docx"},
        }
    )
    assert conflicting["ok"] is False
    assert "conflicting arbitration fee statements" in conflicting["stderr"]

    automatic = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "content": (
                        "# 劳动仲裁申请书\n\n"
                        "## 一、仲裁请求\n\n"
                        "申请人请求支付拖欠工资。"
                    )
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "automatic-legal.docx"},
        }
    )
    assert automatic["ok"] is True, automatic
    automatic_document = Document(
        BytesIO(base64.b64decode(automatic["artifact"]["content_base64"]))
    )
    assert automatic_document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert automatic_document.styles["Normal"].font.size.pt == 12

    reference = Document()
    reference.sections[0].top_margin = Pt(3 * 72 / 2.54)
    reference_header = reference.sections[0].header.paragraphs[0]
    reference_header.text = "参考模板页眉"
    reference_footer = reference.sections[0].footer.paragraphs[0]
    reference_footer.text = "参考模板页脚"
    reference_title = reference.add_paragraph()
    reference_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    reference_title_run = reference_title.add_run("参考标题")
    reference_title_run.font.name = "Courier New"
    reference_title_run.font.size = Pt(20)
    reference_body = reference.add_paragraph()
    reference_body.paragraph_format.first_line_indent = Pt(36)
    reference_body.paragraph_format.line_spacing = 1.5
    reference_body_run = reference_body.add_run("参考正文")
    reference_body_run.font.name = "Courier New"
    reference_body_run.font.size = Pt(13)
    reference_signature = reference.add_paragraph()
    reference_signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    reference_signature.add_run("参考落款")
    reference_stream = BytesIO()
    reference.save(reference_stream)
    templated = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "content": (
                        "# 劳动仲裁申请书\n\n"
                        "正文内容。\n\n"
                        "> 申请人：张三"
                    ),
                    "reference_docx_base64": base64.b64encode(
                        reference_stream.getvalue()
                    ).decode("ascii"),
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "templated-legal.docx"},
        }
    )
    assert templated["ok"] is True, templated
    templated_document = Document(
        BytesIO(base64.b64decode(templated["artifact"]["content_base64"]))
    )
    assert templated_document.sections[0].top_margin == reference.sections[0].top_margin
    assert templated_document.sections[0].header.paragraphs[0].text == "参考模板页眉"
    assert templated_document.sections[0].footer.paragraphs[0].text == "参考模板页脚"
    assert templated_document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert templated_document.paragraphs[0].runs[0].font.name == "Courier New"
    assert templated_document.paragraphs[1].paragraph_format.first_line_indent.pt == 36
    assert templated_document.paragraphs[1].runs[0].font.size.pt == 13
    assert templated_document.paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    invalid_reference = execute_request(
        {
            "skill": "documents",
            "stdin": json.dumps(
                {
                    "content": "# 文档\n\n正文",
                    "reference_docx_base64": "not-base64",
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "docx", "filename": "invalid-reference.docx"},
        }
    )
    assert invalid_reference["ok"] is False
    assert "valid base64" in invalid_reference["stderr"]

    pdf = execute_request(
        {
            "skill": "pdf",
            "stdin": json.dumps(
                {
                    "content": "# 中文报告\n\n"
                    + "\n".join(f"- 第 {index} 项：{long_text}" for index in range(1, 45))
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pdf", "filename": "quality.pdf"},
        }
    )
    assert pdf["ok"] is True, pdf
    import pymupdf

    with pymupdf.open(
        stream=base64.b64decode(pdf["artifact"]["content_base64"]), filetype="pdf"
    ) as pdf_document:
        pdf_text = "\n".join(page.get_text() for page in pdf_document)
        assert "中" in pdf_text
        assert "1" in pdf_text
        assert pdf_document.page_count > 1

    spreadsheet = execute_request(
        {
            "skill": "spreadsheets",
            "stdin": json.dumps(
                {
                    "workbook": {
                        "sheets": [
                            {
                                "name": "Summary",
                                "rows": [
                                    ["Metric", "Value"],
                                    [long_text, 12],
                                    ["第一行\n第二行\n第三行", 3],
                                ],
                            }
                        ]
                    }
                }
            ),
            "artifact": {"format": "xlsx", "filename": "quality.xlsx"},
        }
    )
    assert spreadsheet["ok"] is True, spreadsheet
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(base64.b64decode(spreadsheet["artifact"]["content_base64"])),
        data_only=False,
    )
    sheet = workbook["Summary"]
    assert sheet["A1"].alignment.horizontal == "center"
    assert sheet["B2"].alignment.horizontal == "right"
    assert sheet.sheet_view.showGridLines is False
    assert (sheet.row_dimensions[2].height or 0) > 20
    assert (sheet.row_dimensions[3].height or 0) >= 48
    workbook.close()

    overflowing_bullets = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Quality",
                        "slides": [
                            {
                                "layout": "two_column",
                                "title": "内容必须适合版面",
                                "left": {
                                    "heading": "左侧",
                                    "bullets": ["\n".join(["中" * 12] * 4)] * 3,
                                },
                                "right": {
                                    "heading": "右侧",
                                    "bullets": ["短"] * 3,
                                },
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "overflow.pptx"},
        }
    )
    assert overflowing_bullets["ok"] is False
    assert "does not fit" in overflowing_bullets["stderr"]

    math_slide = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Math",
                        "slides": [
                            {
                                "layout": "bullets",
                                "title": "145 × 12：先估算，再计算",
                                "bullets": [
                                    "145×2＝290",
                                    "145×10＝1450",
                                    "290＋1450＝1740",
                                ],
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "math.pptx"},
        }
    )
    assert math_slide["ok"] is True, math_slide
    with ZipFile(
        BytesIO(base64.b64decode(math_slide["artifact"]["content_base64"]))
    ) as archive:
        slide_xml = archive.read("ppt/slides/slide2.xml").decode()
        theme_xml = archive.read("ppt/theme/theme1.xml").decode()
        assert "算式" in slide_xml
        assert "1740" in slide_xml
        if not os.environ.get("NEXAFLOW_CJK_FONT"):
            pptx_cjk_font = (
                "Noto Serif CJK SC" if sys.platform == "linux" else "Songti SC"
            )
            assert pptx_cjk_font in slide_xml
            assert pptx_cjk_font in theme_xml
