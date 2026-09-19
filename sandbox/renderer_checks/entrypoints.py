"""Preserved renderer regression checks, independent of the retired Broker."""

import base64
import json
from io import BytesIO
from zipfile import ZipFile

try:
    from ..job import execute as execute_request
except ImportError:
    from job import execute as execute_request


def check_builtin_skill_entrypoints() -> None:
    requests = [
        (
            "documents",
            "docx",
            "skill-report.docx",
            {"content": "# Quarterly report\n\nReady for review.\n\n- Revenue\n- Cost"},
            "word/document.xml",
        ),
        (
            "pdf",
            "pdf",
            "skill-report.pdf",
            {"content": "# Quarterly report\n\nReady for review.\n\n- Revenue\n- Cost"},
            None,
        ),
        (
            "pptx",
            "pptx",
            "skill-deck.pptx",
            {
                "presentation": {
                    "title": "Quarterly review",
                    "subtitle": "Decisions and next steps",
                    "template": "bold",
                    "brand": {
                        "primary_color": "#F59E0B",
                        "background_color": "#111827",
                        "text_color": "#F9FAFB",
                        "font_family": "Arial",
                    },
                    "footer": "NexaFlow",
                    "slides": [
                        {
                            "layout": "icons",
                            "title": "Three quarter priorities",
                            "items": [
                                {
                                    "icon": "growth",
                                    "title": "Grow",
                                    "body": "Expand the strongest customer segment.",
                                },
                                {
                                    "icon": "gear",
                                    "title": "Simplify",
                                    "body": "Remove manual steps from delivery.",
                                },
                                {
                                    "icon": "focus",
                                    "title": "Focus",
                                    "body": "Fund the work tied to retention.",
                                },
                            ],
                            "notes": "[Sources]\n- Internal quarterly review",
                        },
                        {
                            "layout": "table",
                            "title": "One outcome per priority",
                            "table": {
                                "headers": ["Priority", "Outcome", "Owner"],
                                "rows": [
                                    ["Grow", "+12% qualified pipeline", "Sales"],
                                    ["Simplify", "-20% delivery time", "Operations"],
                                    ["Focus", "+5 pts retention", "Product"],
                                ],
                            },
                        },
                    ],
                }
            },
            "ppt/presentation.xml",
        ),
        (
            "spreadsheets",
            "xlsx",
            "skill-report.xlsx",
            {
                "workbook": {
                    "sheets": [
                        {
                            "name": "Summary",
                            "rows": [
                                ["Metric", "Value"],
                                ["Revenue", 12],
                                ["Cost", 5],
                                ["Profit", "=B2-B3"],
                            ],
                        }
                    ]
                }
            },
            "xl/workbook.xml",
        ),
    ]
    for skill, artifact_format, filename, inputs, archive_member in requests:
        result = execute_request(
            {
                "skill": skill,
                "stdin": json.dumps(inputs),
                "artifact": {"format": artifact_format, "filename": filename},
            }
        )
        assert result["ok"] is True, result
        content = base64.b64decode(result["artifact"]["content_base64"])
        assert f'"renderer":"{skill}"' in result["stdout"], result
        if artifact_format == "pdf":
            assert content.startswith(b"%PDF"), content[:16]
        else:
            with ZipFile(BytesIO(content)) as archive:
                assert archive_member in archive.namelist(), archive.namelist()

    long_title = "A" * 43
    result = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "bullets",
                                "title": long_title,
                                "bullets": ["Short point"],
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "invalid.pptx"},
        }
    )
    assert result["ok"] is True, result
    missing_table = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "table",
                                "title": "Missing table",
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "missing-table.pptx"},
        }
    )
    assert missing_table["ok"] is False
    assert (
        "presentation.slides[0].table is required for the table layout"
        in missing_table["stderr"]
    )

    long_cell_table = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "table",
                                "title": "Long cells wrap",
                                "table": {
                                    "headers": ["违法情形", "处罚措施"],
                                    "rows": [
                                        [
                                            "擅自举办培训机构（有场所、2名以上人员、有组织机构）",
                                            "责令停止举办、退还费用，处违法所得1—5倍罚款",
                                        ]
                                    ],
                                },
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "long-cells.pptx"},
        }
    )
    assert long_cell_table["ok"] is True, long_cell_table

    two_line_columns = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "two_column",
                                "title": "Two-line bullets",
                                "left": {
                                    "heading": "从轻、减轻或不予处罚",
                                    "bullets": [
                                        "主动消除或减轻危害后果",
                                        "受胁迫、诱骗实施；主动供述未掌握的违法行为",
                                        "配合查处有立功表现",
                                        "轻微并及时改正、无危害后果；能证明无主观过错",
                                    ],
                                },
                                "right": {
                                    "heading": "应当从重处罚",
                                    "bullets": [
                                        "被处理后两年内再次实施同类违法行为",
                                        "危害后果严重、造成恶劣社会影响",
                                        "伪造、涂改、转移、销毁证据",
                                        "拒绝、阻碍或以暴力威胁执法",
                                        "中小学在职教师从事学科类培训",
                                    ],
                                },
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "two-line-columns.pptx"},
        }
    )
    assert two_line_columns["ok"] is True, two_line_columns

    long_title = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "bullets",
                                "title": "结" * 23,
                                "bullets": ["Short point"],
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "long-title.pptx"},
        }
    )
    assert long_title["ok"] is True, long_title

    three_line_subtitle = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "section",
                                "title": "S",
                                "subtitle": "落" * 58,
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "three-line-subtitle.pptx"},
        }
    )
    assert three_line_subtitle["ok"] is True, three_line_subtitle

    two_line_icon_titles = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "icons",
                                "title": "S",
                                "items": [
                                    {
                                        "icon": "star",
                                        "title": "罚款、没收违法所得",
                                        "body": "经济制裁，追缴违法收益",
                                    },
                                    {
                                        "icon": "gear",
                                        "title": "责令停止招收学员、停止举办",
                                        "body": "限制或终止办学行为",
                                    },
                                    {
                                        "icon": "cycle",
                                        "title": "吊销许可证件、限制从业并同步追究相关责任人员责任",
                                        "body": "最严厉处理，并及于责任人员",
                                    },
                                    {
                                        "icon": "heart",
                                        "title": "警告、通报批评",
                                        "body": "最轻处罚",
                                    },
                                ],
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "two-line-icon-titles.pptx"},
        }
    )
    assert two_line_icon_titles["ok"] is True, two_line_icon_titles

    long_stat_value = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "stats",
                                "title": "Long metrics fit on the first render",
                                "stats": [
                                    {"value": "1", "label": "Baseline"},
                                    {
                                        "value": "12345678901234567890",
                                        "label": "Long identifier",
                                    },
                                    {"value": "34.7%", "label": "Growth"},
                                    {"value": "4", "label": "Regions"},
                                ],
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "long-stat.pptx"},
        }
    )
    assert long_stat_value["ok"] is True, long_stat_value

    five_step_titles = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Validation",
                        "slides": [
                            {
                                "layout": "steps",
                                "title": "S",
                                "steps": [
                                    {
                                        "title": "立案与调查",
                                        "body": "符合条件予以立案；现场调查、询问、查阅复制资料",
                                    },
                                    {
                                        "title": "告知与申辩",
                                        "body": "书面告知拟处罚内容及依据，听取陈述申辩并复核",
                                    },
                                    {
                                        "title": "听证与法制审核",
                                        "body": "重大罚款、没收及吊销等案件告知听证权利",
                                    },
                                    {
                                        "title": "决定与送达",
                                        "body": "制作行政处罚决定书，当场交付或七日内送达",
                                    },
                                    {
                                        "title": "执行与结案",
                                        "body": "不履行的申请法院强制执行，执行完毕结案归档",
                                    },
                                ],
                            }
                        ],
                    }
                },
                ensure_ascii=False,
            ),
            "artifact": {"format": "pptx", "filename": "five-step-titles.pptx"},
        }
    )
    assert five_step_titles["ok"] is True, five_step_titles

    custom_theme = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Model-created visual identity",
                        "subtitle": "The renderer protects readability and fit",
                        "theme": {
                            "background_color": "#07111F",
                            "text_color": "#F8FAFC",
                            "accent_color": "#F59E0B",
                            "panel_color": "#172033",
                            "muted_text_color": "#B8C4D6",
                            "rule_color": "#43516A",
                            "font_family": "Arial",
                            "heading_font_family": "Georgia",
                            "cover_title_size": 60,
                            "slide_title_size": 40,
                            "body_size": 20,
                            "panel_radius": 0.12,
                            "cover_accent_width": 0.42,
                            "title_alignment": "center",
                        },
                        "slides": [
                            {
                                "layout": "stats",
                                "title": "One theme, safe geometry",
                                "stats": [
                                    {"value": "1×", "label": "Tool call"},
                                    {"value": "4.5:1", "label": "Text contrast"},
                                    {"value": "35pt", "label": "Title floor"},
                                ],
                            },
                            {
                                "layout": "quote",
                                "title": "A deliberate chapter break",
                                "quote": "Style is model-authored; layout safety stays deterministic.",
                                "source": "NexaFlow renderer contract",
                                "style": {
                                    "background_color": "#FFF8E7",
                                    "text_color": "#2A1A14",
                                    "accent_color": "#8B1E3F",
                                    "panel_color": "#F2E4CC",
                                    "muted_text_color": "#6B4F44",
                                    "rule_color": "#A88C7D",
                                    "title_alignment": "left",
                                    "panel_radius": 0,
                                },
                            },
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "custom-theme.pptx"},
        }
    )
    assert custom_theme["ok"] is True, custom_theme
    from pptx import Presentation
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE

    themed_deck = Presentation(
        BytesIO(base64.b64decode(custom_theme["artifact"]["content_base64"]))
    )
    assert str(themed_deck.slides[0].background.fill.fore_color.rgb) == "07111F"
    assert str(themed_deck.slides[2].background.fill.fore_color.rgb) == "FFF8E7"
    assert any(
        run.font.name == "Georgia"
        for slide in themed_deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    )
    assert any(
        shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        for shape in themed_deck.slides[1].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    )

    partial_theme = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Partial theme",
                        "theme": {
                            "background_color": "#07111F",
                            "text_color": "#F8FAFC",
                        },
                        "slides": [
                            {
                                "layout": "two_column",
                                "title": "Missing tokens get safe defaults",
                                "left": {"heading": "Model", "bullets": ["Theme intent"]},
                                "right": {"heading": "Renderer", "bullets": ["Safe fill"]},
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "partial-theme.pptx"},
        }
    )
    assert partial_theme["ok"] is True, partial_theme

    low_contrast_theme = execute_request(
        {
            "skill": "pptx",
            "stdin": json.dumps(
                {
                    "presentation": {
                        "title": "Invalid contrast",
                        "theme": {
                            "background_color": "#FFFFFF",
                            "text_color": "#F8FAFC",
                        },
                        "slides": [
                            {
                                "layout": "hero",
                                "title": "This must be rejected",
                            }
                        ],
                    }
                }
            ),
            "artifact": {"format": "pptx", "filename": "low-contrast.pptx"},
        }
    )
    assert low_contrast_theme["ok"] is False
    assert "4.5:1 contrast" in low_contrast_theme["stderr"]
